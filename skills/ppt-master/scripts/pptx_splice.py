#!/usr/bin/env python3
"""
PPT Master - pptx_splice

Splice slides from one or more donor PPTX files into a base PPTX in a specified
order, producing a new PPTX. Unreferenced base slides are dropped; donor slides
are copied as native DrawingML shapes (pictures re-embedded). This is the
surgical-edit tool for the edit-pptx workflow: keep most of an existing deck,
swap in a few newly-generated pages, without regenerating the whole deck.

The final slide list is built from an explicit sequence plan (JSON), so delete /
insert / reorder collapse into one operation. Donor slides are added to the
intact base first (new slide parts get names past the base max), then the
sldIdLst is rebuilt in sequence order - this avoids the part-name collisions
that a delete-then-add sequence triggers and that corrupt the output zip.

Usage:
    python3 scripts/pptx_splice.py <plan.json> [--validate]

Examples:
    python3 scripts/pptx_splice.py projects/p/analysis/splice_plan.json --validate

Plan JSON shape:
    {
      "base": "projects/p/sources/deck.pptx",
      "output": "projects/p/exports/deck_edited.pptx",
      "donors": { "new": "projects/p/exports/new_pages.pptx" },
      "sequence": [
        {"src": "base", "slide": 1},
        {"src": "base", "slide": 8},
        {"src": "new",  "slide": 1},
        {"src": "new",  "slide": 2},
        {"src": "base", "slide": 14},
        {"src": "new",  "slide": 5},
        {"src": "base", "slide": 16}
      ]
    }

    src "base"        -> slide N (1-indexed) of the base deck.
    src "<donor>"     -> slide N of that donor file.
    Base slides not listed are dropped. Each listed donor slide is copied once.

Dependencies:
    python-pptx (required)

See workflows/edit-pptx.md for when to use this and the full surgical-edit
workflow. Limitations: native chart/table objects (svg_to_pptx --native-objects),
SmartArt, and hyperlinks are not spliced - flattened shapes and pictures only.
"""

import argparse
import json
import sys
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Optional

try:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError:
    print("python-pptx is required: pip install python-pptx", file=sys.stderr)
    raise

_SCRIPTS_DIR = Path(__file__).resolve().parent
if str(_SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS_DIR))

try:
    from console_encoding import configure_utf8_stdio  # noqa: E402
except ImportError:
    def configure_utf8_stdio() -> None:
        return None


_SHAPE_TAGS = {
    qn("p:sp"), qn("p:pic"), qn("p:grpSp"),
    qn("p:graphicFrame"), qn("p:cxnSp"),
}


def _blank_layout(prs) -> object:
    """Pick a layout with no placeholders (blank); fall back to the last."""
    for layout in prs.slide_layouts:
        try:
            if len(list(layout.placeholders)) == 0:
                return layout
        except Exception:
            continue
    return prs.slide_layouts[-1]


def _add_image_to_slide(dest_slide, src_image_part) -> Optional[str]:
    """Re-embed a source image part into the destination slide's package and
    return the new rId, or None on failure (best-effort picture splicing)."""
    pkg = dest_slide.part.package
    blob = src_image_part.blob
    getter = getattr(pkg, "get_or_add_image_part", None)
    if getter is None:
        return None
    try:
        image_part = getter(BytesIO(blob))
        return dest_slide.part.relate_to(image_part, RT.IMAGE)
    except Exception as e:
        print(f"  warn: image re-embed failed: {e}", file=sys.stderr)
        return None


def _remap_embeds(elem, src_slide, dest_slide) -> None:
    """Re-point r:embed / r:link image references in a copied shape tree from
    the source slide's picture parts to freshly added parts on the destination."""
    src_rels = src_slide.part.rels
    for attr in (qn("r:embed"), qn("r:link")):
        for node in elem.iter(attr):
            old_rid = node.get(attr)
            if not old_rid:
                continue
            rel = src_rels.get(old_rid)
            if rel is None:
                continue
            target = rel.target_part
            if target is None:
                continue
            new_rid = _add_image_to_slide(dest_slide, target)
            if new_rid:
                node.set(attr, new_rid)


def _copy_slide_content(src_slide, dest_slide) -> None:
    """Copy background + shape tree from src_slide into dest_slide (clearing
    dest's default shapes). Picture rIds are remapped per shape."""
    src_cSld = src_slide.element.find(qn("p:cSld"))
    dest_cSld = dest_slide.element.find(qn("p:cSld"))
    if src_cSld is None or dest_cSld is None:
        return
    src_bg = src_cSld.find(qn("p:bg"))
    if src_bg is not None:
        dest_bg = dest_cSld.find(qn("p:bg"))
        if dest_bg is not None:
            dest_cSld.remove(dest_bg)
        dest_cSld.insert(0, deepcopy(src_bg))
    dest_spTree = dest_cSld.find(qn("p:spTree"))
    for child in list(dest_spTree):
        if child.tag in _SHAPE_TAGS:
            dest_spTree.remove(child)
    src_spTree = src_cSld.find(qn("p:spTree"))
    if src_spTree is None:
        return
    for child in list(src_spTree):
        if child.tag not in _SHAPE_TAGS:
            continue
        new_el = deepcopy(child)
        _remap_embeds(new_el, src_slide, dest_slide)
        dest_spTree.append(new_el)


def _related_part(pres_part, rId: str):
    """Resolve a presentation-level rId to its target Part, or None."""
    rel = pres_part.rels.get(rId)
    return rel.target_part if rel is not None else None


def _drop_orphan_slide_parts(pkg, keep_partnames: set) -> None:
    """Best-effort removal of /ppt/slides/slideN.xml parts no longer referenced
    by sldIdLst. Orphan notesSlides are left in place (harmless, no collision)."""
    parts = getattr(pkg, "_parts", None)
    if not isinstance(parts, dict):
        return
    for pname in list(parts.keys()):
        s = str(pname)
        if s.startswith("/ppt/slides/slide") and s.endswith(".xml"):
            if s not in keep_partnames:
                parts.pop(pname, None)


def _validate_no_duplicates(path: str) -> None:
    """Open the saved pptx as a zip and fail if any entry name repeats."""
    with zipfile.ZipFile(path, "r") as z:
        names = z.namelist()
    dups = {n for n in names if names.count(n) > 1}
    if dups:
        raise RuntimeError(
            f"output pptx has duplicate zip entries (corrupt): {sorted(dups)}"
        )


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Splice donor slides into a base PPTX per a sequence plan.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("plan", help="path to the splice plan JSON file")
    p.add_argument(
        "--validate", action="store_true",
        help="after saving, open the zip and assert no duplicate entries",
    )
    return p


def run(plan_path: str, validate: bool) -> str:
    plan = json.loads(Path(plan_path).read_text(encoding="utf-8"))
    base_path = plan["base"]
    out_path = plan["output"]
    donors = plan.get("donors", {})
    sequence = plan["sequence"]

    dest = Presentation(base_path)
    sldIdLst = dest.slides._sldIdLst
    base_ids = list(sldIdLst)
    base_count = len(base_ids)
    donor_prs = {name: Presentation(path) for name, path in donors.items()}

    # 1. Add each unique donor slide once, while the base is still intact.
    #    New slide parts get names past the base max -> no collision with kept
    #    base parts. This is the fix for the delete-then-add collision bug.
    added: dict = {}
    for entry in sequence:
        if entry["src"] == "base":
            continue
        key = (entry["src"], entry["slide"])
        if key in added:
            continue
        dprs = donor_prs[entry["src"]]
        src_slide = dprs.slides[entry["slide"] - 1]
        new_slide = dest.slides.add_slide(_blank_layout(dest))
        _copy_slide_content(src_slide, new_slide)
        added[key] = list(sldIdLst)[-1]

    # 2. Rebuild sldIdLst in the declared sequence order.
    final = []
    for entry in sequence:
        if entry["src"] == "base":
            idx = entry["slide"] - 1
            if idx < 0 or idx >= base_count:
                raise RuntimeError(
                    f"base slide {entry['slide']} out of range "
                    f"(base has {base_count})"
                )
            final.append(base_ids[idx])
        else:
            final.append(added[(entry["src"], entry["slide"])])
    for x in list(sldIdLst):
        sldIdLst.remove(x)
    for x in final:
        sldIdLst.append(x)

    # No manual orphan-part removal: popping _parts directly destabilizes
    # python-pptx's part numbering on save (it reassigns partnames and can mix
    # slide content). Dropped base slides stay in the package as unreferenced
    # parts - harmless (they are not in sldIdLst, so PowerPoint never shows
    # them, and they cannot collide with donor parts added past the base max).
    # --validate is the integrity guard.

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    dest.save(out_path)
    if validate:
        _validate_no_duplicates(out_path)

    dropped = base_count - sum(1 for e in sequence if e["src"] == "base")
    print(
        f"base={base_count} donor_added={len(added)} "
        f"final={len(final)} dropped_base={dropped}",
        file=sys.stderr,
    )
    return out_path


def main(argv: Optional[list] = None) -> int:
    configure_utf8_stdio()
    args = build_parser().parse_args(argv)
    try:
        out = run(args.plan, args.validate)
    except Exception as e:
        print(f"splice failed: {e}", file=sys.stderr)
        return 1
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
